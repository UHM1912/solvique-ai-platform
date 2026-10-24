import json
import uuid
from io import BytesIO
from pathlib import Path
from typing import Any

import plotly.graph_objects as go
import plotly.io as pio
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from config import GCS_BUCKET_NAME, REPORT_ASSET_BACKEND, REPORT_ASSET_ROOT
from src.conversation_store import add_report_asset
from src.schemas.api_models import ReportSection


class LocalReportAssetStorage:
    def __init__(self, root: Path | None = None):
        self.root = Path(root or REPORT_ASSET_ROOT)
        self.root.mkdir(parents=True, exist_ok=True)

    def save_bytes(self, data: bytes, object_path: str) -> str:
        object_path = _validated_object_path(object_path)
        target = (self.root / object_path).resolve()
        if not target.is_relative_to(self.root.resolve()):
            raise ValueError("Invalid report asset path")
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(data)
        return object_path

    def resolve_path(self, object_path: str) -> Path:
        path = (self.root / _validated_object_path(object_path)).resolve()
        if not path.is_relative_to(self.root.resolve()):
            raise ValueError("Invalid report asset path")
        return path

    def read_bytes(self, object_path: str) -> bytes:
        return self.resolve_path(object_path).read_bytes()

    def delete_object(self, object_path: str) -> None:
        self.resolve_path(object_path).unlink(missing_ok=True)


def _validated_object_path(object_path: str) -> str:
    value = object_path.replace("\\", "/")
    if not value or value.startswith("/") or any(part in {"", ".", ".."} for part in value.split("/")):
        raise ValueError("Invalid report asset path")
    return value


class GCSReportAssetStorage:
    """Private GCS object storage authenticated with Application Default Credentials."""

    def __init__(self, bucket_name: str):
        if not bucket_name:
            raise RuntimeError("GCS report storage requires GCS_BUCKET_NAME")
        try:
            from google.cloud import storage
        except ImportError as exc:
            raise RuntimeError("GCS report storage requires google-cloud-storage") from exc
        self.bucket = storage.Client().bucket(bucket_name)

    def save_bytes(self, data: bytes, object_path: str) -> str:
        path = _validated_object_path(object_path)
        self.bucket.blob(path).upload_from_string(data, content_type="image/png")
        return path

    def read_bytes(self, object_path: str) -> bytes:
        from google.api_core.exceptions import NotFound
        try:
            return self.bucket.blob(_validated_object_path(object_path)).download_as_bytes()
        except NotFound as exc:
            raise FileNotFoundError(object_path) from exc

    def resolve_path(self, object_path: str) -> Path:
        # Kept for callers that require a filesystem path. Report rendering uses
        # read_bytes so cloud assets do not need persistent local copies.
        suffix = Path(_validated_object_path(object_path)).suffix or ".bin"
        from tempfile import NamedTemporaryFile
        with NamedTemporaryFile(prefix="solvique-report-", suffix=suffix, delete=False) as tmp:
            tmp.write(self.read_bytes(object_path))
            return Path(tmp.name)

    def delete_object(self, object_path: str) -> None:
        from google.api_core.exceptions import NotFound
        try:
            self.bucket.blob(_validated_object_path(object_path)).delete()
        except NotFound:
            pass


class ChartAssetService:
    def __init__(self, storage: LocalReportAssetStorage):
        self.storage = storage

    def create_chart_asset(
        self,
        *,
        conversation_id: str,
        message_id: str,
        figure_json: str,
    ) -> dict[str, Any]:
        asset_id = f"asset_{uuid.uuid4().hex}"
        object_path = self.storage.save_bytes(
            data=pio.to_image(
                self.figure(figure_json), format="png", width=1400, height=900, scale=2
            ),
            object_path=f"{conversation_id}/{message_id}/{asset_id}.png",
        )
        return add_report_asset(
            id=asset_id,
            conversation_id=conversation_id,
            message_id=message_id,
            object_path=object_path,
        )

    def figure(self, figure_json: str) -> go.Figure:
        try:
            return pio.from_json(figure_json, skip_invalid=True)
        except TypeError:
            return go.Figure(json.loads(figure_json), skip_invalid=True)


class ReportDocxRenderer:
    def __init__(self, asset_path_resolver=None):
        self.asset_path_resolver = asset_path_resolver

    def render(self, sections: list[ReportSection]) -> bytes:
        doc = Document()
        self.configure(doc)

        doc.add_heading("Analysis", level=1)
        analysis_index = 0

        for section in [item for item in sections if item.kind != "appendix"]:
            if section.kind == "analysis":
                analysis_index += 1
                doc.add_heading(f"{analysis_index}. {section.heading}", level=2)
                self.add_findings(doc, section.body)
            elif section.kind == "visualization":
                doc.add_heading(section.heading, level=3)
                self.add_note(doc, section.body)
                self.add_chart(doc, section.asset_id)
            elif section.kind == "table":
                doc.add_heading(section.heading, level=3)
                self.add_note(doc, section.body)
                self.add_data_table(doc, section.table_columns, section.table_rows)

        appendix = [item for item in sections if item.kind == "appendix"]
        if appendix:
            doc.add_page_break()
            doc.add_heading("Appendix", level=1)
            for section in appendix:
                doc.add_heading(section.heading, level=2)
                self.add_code_block(doc, section.body)

        buffer = BytesIO()
        doc.save(buffer)
        return buffer.getvalue()

    def configure(self, doc: Document) -> None:
        page = doc.sections[0]
        page.top_margin = page.bottom_margin = Inches(0.7)
        page.left_margin = page.right_margin = Inches(0.75)

        doc.styles["Normal"].font.name = "Calibri"
        doc.styles["Normal"].font.size = Pt(10.5)
        doc.styles["Normal"].paragraph_format.space_after = Pt(6)
        for name, size, color in (
            ("Heading 1", 16, "0F172A"),
            ("Heading 2", 12.5, "111827"),
            ("Heading 3", 11, "0F766E"),
        ):
            style = doc.styles[name]
            style.font.name = "Calibri"
            style.font.bold = True
            style.font.size = Pt(size)
            style.font.color.rgb = RGBColor.from_string(color)
            style.paragraph_format.space_before = Pt(8 if name == "Heading 2" else 6)
            style.paragraph_format.space_after = Pt(7 if name == "Heading 2" else 5)

    def add_findings(self, doc: Document, body: str) -> None:
        cell = self.one_cell_table(doc, fill="F8FAFC", accent="0F766E")
        title = cell.paragraphs[0]
        title.paragraph_format.space_after = Pt(8)
        run = title.add_run("Key findings")
        run.bold = True
        run.font.size = Pt(10.5)

        for item in self.findings(body):
            paragraph = cell.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.left_indent = Inches(0.28)
            paragraph.paragraph_format.first_line_indent = Inches(-0.18)
            paragraph.paragraph_format.space_before = Pt(2)
            paragraph.paragraph_format.space_after = Pt(5)
            paragraph.paragraph_format.line_spacing = 1.08
            paragraph.add_run(item)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def add_chart(self, doc: Document, asset_id: str | None) -> None:
        if not asset_id or not self.asset_path_resolver:
            return

        resolved = self.asset_path_resolver(asset_id)
        if resolved is None:
            return
        if isinstance(resolved, (bytes, bytearray)):
            image_source = BytesIO(resolved)
        elif hasattr(resolved, "read"):
            image_source = resolved
        else:
            path = Path(resolved)
            if not path.exists():
                return
            image_source = str(path)

        cell = self.one_cell_table(doc)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(image_source, width=Inches(6.2))
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    def add_data_table(
        self, doc: Document, columns: list[str], rows: list[dict[str, Any]]
    ) -> None:
        if not columns or not rows:
            return

        columns = columns[:6]
        table = doc.add_table(rows=1, cols=len(columns))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"

        for index, column in enumerate(columns):
            cell = table.rows[0].cells[index]
            self.shade(cell, "E2E8F0")
            self.cell_text(cell, column.replace("_", " ").title(), bold=True)

        for row in rows:
            cells = table.add_row().cells
            for index, column in enumerate(columns):
                self.cell_text(cells[index], self.truncate(row.get(column)))

        for row in table.rows:
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    def add_note(self, doc: Document, text: str) -> None:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(5)
        run = paragraph.add_run(text)
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor.from_string("475569")

    def add_code_block(self, doc: Document, text: str) -> None:
        for chunk in self.chunks(text, 1800):
            cell = self.one_cell_table(doc, fill="F8FAFC")
            run = cell.paragraphs[0].add_run(chunk)
            run.font.name = "Courier New"
            run.font.size = Pt(8.5)

    def one_cell_table(
        self, doc: Document, *, fill: str = "FFFFFF", accent: str | None = None
    ):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        self.shade(cell, fill)
        self.border(cell, color="CBD5E1", accent=accent)
        self.cell_margin(cell, top=120, left=180, bottom=120, right=180)
        return cell

    def findings(self, body: str) -> list[str]:
        items = [
            line[2:].strip()
            for line in body.splitlines()
            if line.strip().startswith("- ")
        ]
        return items or [body.replace("Key findings:", "").strip()]

    def cell_text(self, cell, text: str, *, bold: bool = False) -> None:
        cell.text = ""
        run = cell.paragraphs[0].add_run(str(text))
        run.bold = bold
        run.font.size = Pt(8.8)

    def shade(self, cell, fill: str) -> None:
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), fill)
        tc_pr.append(shd)

    def border(self, cell, *, color: str, accent: str | None = None) -> None:
        tc_pr = cell._tc.get_or_add_tcPr()
        borders = OxmlElement("w:tcBorders")
        for edge in ("top", "left", "bottom", "right"):
            tag = OxmlElement(f"w:{edge}")
            tag.set(qn("w:val"), "single")
            tag.set(qn("w:sz"), "12" if accent and edge == "left" else "6")
            tag.set(qn("w:space"), "0")
            tag.set(qn("w:color"), accent if accent and edge == "left" else color)
            borders.append(tag)
        tc_pr.append(borders)

    def cell_margin(
        self, cell, *, top: int, left: int, bottom: int, right: int
    ) -> None:
        tc_pr = cell._tc.get_or_add_tcPr()
        margins = OxmlElement("w:tcMar")
        for edge, value in (
            ("top", top),
            ("left", left),
            ("bottom", bottom),
            ("right", right),
        ):
            tag = OxmlElement(f"w:{edge}")
            tag.set(qn("w:w"), str(value))
            tag.set(qn("w:type"), "dxa")
            margins.append(tag)
        tc_pr.append(margins)

    def chunks(self, text: str, limit: int) -> list[str]:
        value = text or ""
        parts = []
        while len(value) > limit:
            split_at = max(value.rfind("\n", 0, limit), value.rfind(" ", 0, limit))
            split_at = split_at if split_at > limit // 2 else limit
            parts.append(value[:split_at].strip())
            value = value[split_at:].strip()
        return [*parts, value] if value else parts or [""]

    def truncate(self, value: Any) -> str:
        text = "" if value is None else str(value)
        return f"{text[:57]}..." if len(text) > 60 else text


if REPORT_ASSET_BACKEND == "gcs":
    asset_storage = GCSReportAssetStorage(GCS_BUCKET_NAME)
else:
    asset_storage = LocalReportAssetStorage()
chart_asset_service = ChartAssetService(asset_storage)
